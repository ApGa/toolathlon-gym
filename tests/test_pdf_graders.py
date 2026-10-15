"""PDF deliverables must be read, not just found.

The image installs pypdf only. Three graders imported PyPDF2 and then pdfplumber,
found neither and checked only that the PDF existed, so any file of the right
name and size passed. Graders now read PDFs through grader_helpers.pdf_text,
which tries pypdf first, keeps the older readers as fallbacks, and raises when
no reader is installed or none can parse the file; the graders turn that into
a failed check.
"""
import contextlib
import datetime
import importlib.util
import io
from pathlib import Path
import random
import sys
from tempfile import TemporaryDirectory
from types import ModuleType, SimpleNamespace
import unittest
from unittest.mock import patch
import uuid

import psycopg2

ROOT = Path(__file__).resolve().parents[1]
TASKS = ROOT / "tasks" / "finalpool"
sys.path.insert(0, str(ROOT / "runtime_patches"))
from grader_helpers import PdfTextError, pdf_text  # noqa: E402

NO_PDF_READERS = {"pypdf": None, "PyPDF2": None, "pdfplumber": None}


def _escape(text):
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(lines, lines_per_page=50):
    """A text PDF (Helvetica, one text line per entry) like simple report writers produce."""
    pages = [lines[i:i + lines_per_page] for i in range(0, len(lines), lines_per_page)] or [[]]
    objects = {
        1: "<< /Type /Catalog /Pages 2 0 R >>",
        3: "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    }
    kids, number = [], 4
    for page in pages:
        content = "BT /F1 10 Tf 13 TL 50 800 Td " + " ".join(f"({_escape(line)}) Tj T*" for line in page) + " ET"
        objects[number] = f"<< /Length {len(content.encode('latin-1'))} >>\nstream\n{content}\nendstream"
        objects[number + 1] = ("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] "
                               f"/Resources << /Font << /F1 3 0 R >> >> /Contents {number} 0 R >>")
        kids.append(f"{number + 1} 0 R")
        number += 2
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for n in sorted(objects):
        offsets[n] = len(out)
        out += f"{n} 0 obj\n{objects[n]}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {number}\n0000000000 65535 f \n".encode()
    for n in range(1, number):
        out += f"{offsets[n]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {number} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


NOTES = [f"Note {i}: figures come from the systems of record as of the report date." for i in range(12)]

# What each task asks its PDF to contain, written so that every check on it passes.
GOOD_PDFS = {
    "sf-sales-trend-word-pdf": ("Sales_Trend_Analysis.pdf", [
        "Sales Trend Analysis 2025",
        *[f"2025-{m:02d}  revenue ${100000 + 1000 * m:,}.00" for m in range(1, 13)],
        *NOTES,
    ]),
    "wc-inventory-report-pdf-gcal": ("Inventory_Audit.pdf", [
        "Inventory Audit Report",
        "Total products: 82",
        "Out of stock products (5): 20, 39, 45, 71, 79",
        "Low stock products (1-5 units): 5, 7, 15, 19, 21, 22, 27, 28, 30, 31, 32, 50, 52, 53",
        *NOTES,
    ]),
    "yf-dividend-analysis-pdf-email": ("Dividend_Analysis.pdf", [
        "Dividend Analysis Report",
        "Symbol  Sector  Dividend yield",
        "XOM  Energy  2.75%",
        "JNJ  Healthcare  2.12%",
        "JPM  Financial Services  2.00%",
        "GOOGL  Technology  0.28%",
        "Ranking: 1. XOM has the highest dividend yield",
        "Recommendation: hold XOM and JNJ for income.",
        *NOTES,
    ]),
}


def garbage(size=4096):
    return random.Random(7).randbytes(size)


def load_grader(task):
    path = TASKS / task / "evaluation" / "main.py"
    spec = importlib.util.spec_from_file_location(f"grader_{uuid.uuid4().hex}", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def check_lines(output):
    lines = [line.strip() for line in output.splitlines()]
    passed = {line[len("[PASS] "):].split(":")[0] for line in lines if line.startswith("[PASS] ")}
    failed = {line[len("[FAIL] "):].split(":")[0] for line in lines if line.startswith("[FAIL] ")}
    return passed, failed


class PdfTextTest(unittest.TestCase):
    def setUp(self):
        tmp = TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)

    def write(self, name, data):
        path = self.root / name
        path.write_bytes(data)
        return path

    def test_reads_the_text_with_pypdf(self):
        path = self.write("a.pdf", make_pdf(["Inventory Audit", "Out of stock: 20, 39"] * 30, lines_per_page=40))
        text = pdf_text(path)
        self.assertIn("Inventory Audit", text)
        self.assertIn("Out of stock: 20, 39", text)
        self.assertEqual(text.count("Inventory Audit"), 30, "every page is read")

    def test_an_empty_or_garbage_file_raises(self):
        for name, data in (("empty.pdf", b""), ("garbage.pdf", garbage()), ("text.pdf", b"Inventory Audit\n" * 200)):
            with self.subTest(name=name), self.assertRaises(PdfTextError) as error:
                pdf_text(self.write(name, data))
            self.assertIn("pypdf could not parse it", str(error.exception))

    def test_no_installed_reader_raises(self):
        path = self.write("a.pdf", make_pdf(["Inventory Audit"]))
        with patch.dict(sys.modules, NO_PDF_READERS), self.assertRaises(PdfTextError) as error:
            pdf_text(path)
        self.assertIn("pypdf is not installed; PyPDF2 is not installed; pdfplumber is not installed",
                      str(error.exception))

    def test_older_readers_are_fallbacks_in_order(self):
        path = self.write("a.pdf", b"%PDF-1.4 ...")
        page = SimpleNamespace(extract_text=lambda: "from PyPDF2")
        pypdf2 = ModuleType("PyPDF2")
        pypdf2.PdfReader = lambda _path: SimpleNamespace(pages=[page])

        class Plumbed:
            pages = [SimpleNamespace(extract_text=lambda: "from pdfplumber")]

            def __enter__(self):
                return self

            def __exit__(self, *exc):
                return False

        plumber = ModuleType("pdfplumber")
        plumber.open = lambda _path: Plumbed()
        broken = ModuleType("pypdf")

        def cannot_parse(_path):
            raise ValueError("unsupported filter")

        broken.PdfReader = cannot_parse
        with patch.dict(sys.modules, {"pypdf": None, "PyPDF2": pypdf2, "pdfplumber": plumber}):
            self.assertEqual(pdf_text(path), "from PyPDF2")
        with patch.dict(sys.modules, {"pypdf": None, "PyPDF2": None, "pdfplumber": plumber}):
            self.assertEqual(pdf_text(path), "from pdfplumber")
        with patch.dict(sys.modules, {"pypdf": broken, "PyPDF2": None, "pdfplumber": plumber}):
            self.assertEqual(pdf_text(path), "from pdfplumber", "a reader that cannot parse is skipped")

    def test_ligatures_read_as_letters(self):
        # Chromium and office suites print "fi", "fl" and "ff" as ligature glyphs, which pypdf returns as
        # U+FB01, U+FB02 and U+FB00: sf-sales-region-pdf-gsheet's "asia pacific" check failed on a printed PDF.
        path = self.write("a.pdf", b"%PDF-1.4 ...")
        printed = ModuleType("pypdf")
        printed.PdfReader = lambda _path: SimpleNamespace(
            pages=[SimpleNamespace(extract_text=lambda: "Asia Paciﬁc, cash ﬂow, staﬀ")])
        with patch.dict(sys.modules, {"pypdf": printed}):
            self.assertEqual(pdf_text(path), "Asia Pacific, cash flow, staff")


class PdfCheckTest(unittest.TestCase):
    """check_pdf of the three graders the c1e58af review found checking existence only."""

    def setUp(self):
        tmp = TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.workspace = Path(tmp.name)

    def run_check(self, task, data, modules=None):
        name, _ = GOOD_PDFS[task]
        (self.workspace / name).write_bytes(data)
        grader = load_grader(task)
        out = io.StringIO()
        with contextlib.redirect_stdout(out), patch.dict(sys.modules, modules or {}):
            grader.check_pdf(str(self.workspace))
        return grader, out.getvalue()

    def test_a_correct_pdf_passes_every_check(self):
        for task, (_, lines) in GOOD_PDFS.items():
            with self.subTest(task=task):
                data = make_pdf(lines)
                self.assertGreater(len(data), 1024)
                grader, output = self.run_check(task, data)
                self.assertEqual(grader.FAIL_COUNT, 0, output)
                self.assertGreaterEqual(grader.PASS_COUNT, 4, output)

    def test_an_empty_or_garbage_pdf_fails(self):
        for task in GOOD_PDFS:
            for label, data in (("empty", b""), ("garbage", garbage())):
                with self.subTest(task=task, pdf=label):
                    grader, output = self.run_check(task, data)
                    passed, failed = check_lines(output)
                    self.assertIn("PDF text readable", failed, output)
                    self.assertIn("could not parse it", output)
                    # Every content check fails with it, so it scores below a readable wrong PDF.
                    self.assertLessEqual(passed, {"PDF file exists", "PDF file size reasonable"}, output)

    def test_no_pdf_reader_fails_closed(self):
        for task, (_, lines) in GOOD_PDFS.items():
            with self.subTest(task=task):
                grader, output = self.run_check(task, make_pdf(lines), NO_PDF_READERS)
                passed, failed = check_lines(output)
                self.assertIn("PDF text readable", failed, output)
                self.assertIn("pypdf is not installed", output)
                self.assertNotIn("existence only", output)
                self.assertEqual(passed, {"PDF file exists", "PDF file size reasonable"}, output)


def monthly_revenue():
    return [(datetime.date(2025, m, 1), 100000 + 1000 * m, 400 + m, 300 + m, 250.0) for m in range(1, 13)]


def sales_trend_docx(path):
    from docx import Document
    doc = Document()
    doc.add_heading("Sales Trend Analysis 2025", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text, table.rows[0].cells[2].text = "Month", "Revenue", "Orders"
    for month, revenue, orders, _customers, _aov in monthly_revenue():
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = month.strftime("%Y-%m"), f"${revenue:,.2f}", str(orders)
    doc.add_paragraph(f"Total revenue for 2025: ${sum(r[1] for r in monthly_revenue()):,.2f}")
    doc.add_paragraph("H1 versus H2: revenue grew in the second half of the year.")
    doc.save(path)


class PdfGraderVerdictTest(unittest.TestCase):
    """The whole grade: a correct submission passes; a junk PDF or no PDF reader fails it."""

    def grade(self, task, pdf, modules=None):
        with TemporaryDirectory() as tmp:
            workspace = Path(tmp)
            name, _ = GOOD_PDFS[task]
            (workspace / name).write_bytes(pdf)
            if task == "sf-sales-trend-word-pdf":
                sales_trend_docx(workspace / "Sales_Trend_Analysis.docx")
            grader = load_grader(task)
            if hasattr(grader, "get_expected_data"):
                grader.get_expected_data = monthly_revenue
            # The calendar and e-mail checks are non-blocking; no database here.
            refuse = patch.object(grader.psycopg2, "connect",
                                  side_effect=psycopg2.OperationalError("no database in this test"))
            out = io.StringIO()
            argv = ["main.py", "--agent_workspace", str(workspace), "--res_log_file", str(workspace / "res.json")]
            with refuse, patch.object(sys, "argv", argv), contextlib.redirect_stdout(out), \
                    patch.dict(sys.modules, modules or {}), self.assertRaises(SystemExit) as code:
                grader.main()
            return code.exception.code, out.getvalue()

    def test_verdicts(self):
        for task, (_, lines) in GOOD_PDFS.items():
            with self.subTest(task=task, pdf="correct"):
                code, output = self.grade(task, make_pdf(lines))
                self.assertEqual(code, 0, output)
                self.assertIn("Overall: PASS", output)
            with self.subTest(task=task, pdf="garbage"):
                code, output = self.grade(task, garbage())
                self.assertEqual(code, 1, output)
                self.assertIn("[FAIL] PDF text readable", output)
            with self.subTest(task=task, pdf="correct, no reader installed"):
                code, output = self.grade(task, make_pdf(lines), NO_PDF_READERS)
                self.assertEqual(code, 1, output)
                self.assertIn("[FAIL] PDF text readable", output)


class ShoppingListPdfTest(unittest.TestCase):
    """howtocook-workshop-materials credited any PDF over 2 KB when its text was missing or unreadable."""

    def check(self, data, modules=None):
        with TemporaryDirectory() as tmp:
            (Path(tmp) / "Shopping_List.pdf").write_bytes(data)
            grader = load_grader("howtocook-workshop-materials")
            out = io.StringIO()
            with contextlib.redirect_stdout(out), patch.dict(sys.modules, modules or {}):
                passed, total = grader.check_pdf(tmp)
            return passed, total, out.getvalue()

    def test_text_decides_not_size(self):
        shopping = make_pdf(["Shopping List", *[f"- {i * 100} g ingredient {i}" for i in range(60)]])
        unrelated = make_pdf([f"Workshop agenda item {i}" for i in range(80)])
        self.assertGreater(len(unrelated), 2000)
        self.assertEqual(self.check(shopping)[:2], (3, 3))
        for label, data, modules in (("no shopping list text", unrelated, None),
                                     ("garbage", garbage(), None),
                                     ("no reader installed", shopping, NO_PDF_READERS)):
            with self.subTest(pdf=label):
                passed, total, output = self.check(data, modules)
                self.assertEqual((passed, total), (2, 3), output)
                self.assertIn("FAIL", output)


if __name__ == "__main__":
    unittest.main()
