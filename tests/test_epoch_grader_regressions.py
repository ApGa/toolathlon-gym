"""Regression reports use retrieved data, never illustrative reference workbooks."""
import contextlib
import importlib.util
import io
import json
import os
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import openpyxl
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "runtime_patches"))
from grader_helpers import rich_text


def grader(task):
    path = ROOT / "tasks/finalpool" / task / "evaluation/main.py"
    spec = importlib.util.spec_from_file_location("epoch_grader", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def workbook(tables):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in tables.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    return book


class EpochGraderRegressions(unittest.TestCase):
    def run_check(self, module, fn, *args):
        module.PASS_COUNT = module.FAIL_COUNT = 0
        with contextlib.redirect_stdout(io.StringIO()):
            fn(*args)
        return module.FAIL_COUNT

    def test_notion_input_and_output_title_forms_match(self):
        expected = "Conference Reading List"
        for value in [expected, [{"text": {"content": expected}}], [{"plain_text": expected}],
                      json.dumps([{"text": {"content": expected}}])]:
            self.assertEqual(rich_text(value), expected)

    def test_conference_task_does_not_require_unrequested_spreadsheet(self):
        m = grader("fetch-arxiv-conference-schedule-gcal-notion")
        with patch.object(m, "check_notion", return_value=True), patch.object(m, "check_calendar", return_value=True), \
             patch.object(sys, "argv", ["grader"]), contextlib.redirect_stdout(io.StringIO()):
            with self.assertRaises(SystemExit) as result:
                m.main()
        self.assertEqual(result.exception.code, 0)

    def test_conference_failed_individual_check_prevents_success(self):
        m = grader("fetch-arxiv-conference-schedule-gcal-notion")
        m.FAIL_COUNT = 1
        with patch.object(m, "check_notion", return_value=True), patch.object(m, "check_calendar", return_value=True), \
             patch.object(sys, "argv", ["grader"]), contextlib.redirect_stdout(io.StringIO()):
            with self.assertRaises(SystemExit) as result:
                m.main()
        self.assertEqual(result.exception.code, 1)

    def customer_report(self):
        return workbook({
            "Customer_Segments": [["Segment", "Customer_Count", "Total_Revenue", "Avg_Spend", "Avg_Orders"],
                                  ["VIP", 1, 600, 600, 3], ["Regular", 2, 600, 300, 2], ["New", 1, 50, 50, 1]],
            "Top_Customers": [["Customer_Name", "Email", "Total_Spent", "Order_Count", "Segment"],
                              ["Grace Hopper", "grace@store.example", 600, 3, "VIP"],
                              ["Ada Lovelace", "ada@store.example", 500, 2, "Regular"],
                              ["Alan Turing", "alan@store.example", 100, 2, "Regular"],
                              ["Katherine Johnson", "kat@store.example", 50, 1, "New"]],
            "Segment_Strategy": [["Segment", "Engagement_Strategy", "Retention_Risk", "Target_Action"],
                                 ["VIP", "Loyalty", "Low", "Offer"], ["Regular", "Discount", "Medium", "Upsell"],
                                 ["New", "Welcome", "High", "Onboard"]],
        })

    def test_store_customer_report_and_tier_boundaries(self):
        m = grader("wc-customer-gform-excel-notion")
        customers = [("Grace Hopper", "grace@store.example", 600, 3), ("Ada Lovelace", "ada@store.example", 500, 2),
                     ("Alan Turing", "alan@store.example", 100, 2), ("Katherine Johnson", "kat@store.example", 50, 1)]
        book = self.customer_report()
        self.assertEqual(self.run_check(m, m.check_customer_data, book, customers), 0)
        book["Top_Customers"]["B2"] = "a@example.com"
        book["Customer_Segments"]["C2"] = 1000000
        self.assertGreaterEqual(self.run_check(m, m.check_customer_data, book, customers), 2)

    def test_curriculum_uses_real_enrollment_and_paper_metadata(self):
        m = grader("canvas-scholarly-curriculum-excel-notion")
        book = workbook({
            "Current_Courses": [["Course_Name", "Course_Code", "Enrollment_Count", "Avg_Score"],
                                ["Applied Analytics", "AAA-2013J", 383, 65.1]],
            "Research_Trends": [["Paper_Title", "Topic_Area", "Year", "Citations", "Relevance_to_Curriculum"],
                                ["Real Data Analytics Paper", "Data Analytics", 2025, 87, "High"]],
            "Gap_Analysis": [["Metric", "Value"], ["Total_Courses", 1], ["Papers_Reviewed", 1],
                             ["High_Relevance_Papers", 1], ["Curriculum_Coverage_Pct", 50], ["Top_Gap_Area", "Deep Learning"]],
        })
        courses = [("Applied Analytics", "AAA-2013J", 383, 65.123)]
        papers = [("Real Data Analytics Paper", 2025, 87)]
        self.assertEqual(self.run_check(m, m.check_curriculum, book, courses, papers), 0)
        book["Current_Courses"]["C2"] = 500
        book["Research_Trends"]["A2"] = "Deep Learning Advances"
        self.assertGreaterEqual(self.run_check(m, m.check_curriculum, book, courses, papers), 2)

    def test_salary_fixture_and_report_not_sales_placeholder(self):
        m = grader("fetch-sf-sales-competitor-excel-notion")
        benchmarks = m.load_benchmarks()
        self.assertEqual(benchmarks["Engineering"], 66816)
        salaries = {d: b + 100 for d, b in benchmarks.items()}
        book = workbook({
            "Data_Analysis": [["Department", "Internal_Avg_Salary", "Industry_Avg", "Gap"]] +
                             [[d, salaries[d], benchmarks[d], 100] for d in sorted(benchmarks)],
            "Metrics": [["Metric", "Value"], ["Departments", 7], ["Mean Gap", 100], ["Average", 70000], ["Employees", 100]],
            "Recommendations": [["Priority", "Action"], [1, "Review compensation"], [2, "Survey salaries"]],
        })
        self.assertEqual(self.run_check(m, m.check_comparison, book, salaries, benchmarks), 0)
        book["Data_Analysis"]["C2"] = 1
        self.assertGreater(self.run_check(m, m.check_comparison, book, salaries, benchmarks), 0)

    def test_assignment_report_checks_actual_requested_course(self):
        m = grader("canvas-assignment-feedback-gform-word")
        expected = [("CMA 34878", 100, 75.25, 10, 10.0), ("Final Exam 34885", 0, None, 0, 0),
                    ("TMA 34880", 150, 85.5, 5, 3.33)]
        with TemporaryDirectory() as tmp:
            doc = Document()
            doc.add_heading("Foundations of Finance (Fall 2013) - Assignment Analysis", 0)
            table = doc.add_table(rows=1, cols=5)
            for cell, header in zip(table.rows[0].cells, ["Assignment_Name", "Total_Submissions", "Avg_Score", "Late_Submissions", "Late_Rate(%)"]):
                cell.text = header
            for row in expected:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = "N/A" if value is None else str(value)
            doc.save(Path(tmp) / "Assignment_Analysis.docx")
            self.assertEqual(self.run_check(m, m.check_word, tmp, tmp, expected), 0)
            table.rows[1].cells[0].text = "CMA 34905"
            doc.save(Path(tmp) / "Assignment_Analysis.docx")
            self.assertGreater(self.run_check(m, m.check_word, tmp, tmp, expected), 0)

    def test_fireship_reports_seeded_stats_and_topic_aggregates(self):
        m = grader("yt-fireship-gform-survey-excel-gcal")
        references = [(f"video-{i}", f"Seeded title {i}", 8000 - i * 100, 80 - i, "120") for i in range(8)]
        cursor = unittest.mock.MagicMock()
        cursor.__enter__.return_value = cursor
        cursor.fetchall.return_value = references
        conn = SimpleNamespace(cursor=lambda: cursor, close=lambda: None)
        rows = [["Rank", "Video_ID", "Title", "Views", "Likes", "Duration_Sec", "Topic_Tags", "Engagement_Rate"]]
        topic_rows = [["Topic", "Avg_Engagement_Rate", "Total_Views", "Video_Count"]]
        for i, (vid, title, views, likes, duration) in enumerate(references):
            rows.append([i + 1, vid, title, views, likes, int(duration), f"Topic {i // 2}", 100 * likes / views])
        for i in range(4):
            vids = references[2*i:2*i+2]
            topic_rows.append([f"Topic {i}", sum(100 * v[3] / v[2] for v in vids) / 2, sum(v[2] for v in vids), 2])
        book = workbook({"Top_Videos": rows, "Engagement_Analysis": topic_rows})
        with patch.object(m.psycopg2, "connect", return_value=conn):
            self.assertEqual(self.run_check(m, m.check_seeded_video_data, book), 0)
            book["Top_Videos"]["B2"] = "placeholder-video"
            book["Engagement_Analysis"]["C2"] = 1
            self.assertGreater(self.run_check(m, m.check_seeded_video_data, book), 0)

    def test_health_calendar_accepts_different_order_with_equal_scores(self):
        from datetime import datetime
        m = grader("sf-customer-health-dashboard")
        expected = {"below_15": [("Alice", "", "", 0, 0, 0, 5), ("Bob", "", "", 0, 0, 0, 5),
                                 ("Charlie", "", "", 0, 0, 0, 8)]}
        cursor = unittest.mock.MagicMock()
        cursor.fetchall.return_value = [("Follow-up Call: Bob", "", datetime(2026, 3, 9)),
                                        ("Follow-up Call: Alice", "", datetime(2026, 3, 10)),
                                        ("Follow-up Call: Charlie", "", datetime(2026, 3, 11))]
        conn = SimpleNamespace(cursor=lambda: cursor, close=lambda: None)
        with patch.object(m.psycopg2, "connect", return_value=conn):
            self.assertEqual(self.run_check(m, m.check_calendar, expected), 0)
            cursor.fetchall.return_value.reverse()
            self.assertGreater(self.run_check(m, m.check_calendar, expected), 0)

    def test_compliance_notion_accepts_text_content_titles(self):
        m = grader("playwright-canvas-curriculum-word-notion")
        cursor = unittest.mock.MagicMock()
        pages = [(str(i), {"Overall_Status": {"select": {"name": "Non-Compliant" if i < 9 else "Compliant"}}}) for i in range(22)]
        cursor.fetchall.side_effect = [[("db-id", [{"text": {"content": "Course Compliance Tracker"}}], {})], pages]
        conn = SimpleNamespace(cursor=lambda: cursor, close=lambda: None)
        with patch.object(m.psycopg2, "connect", return_value=conn):
            self.assertEqual(self.run_check(m, m.check_notion), 0)

    def test_grade_report_checks_google_sheet_data_without_xlsx(self):
        m = grader("canvas-grades-gsheet-pdf-email")
        grade = {"course_name": "Course A", "a_count": 1, "b_count": 1, "c_count": 1,
                 "d_count": 1, "f_count": 1, "total_students": 5, "pass_rate_pct": 60, "course_avg": 70}
        metrics = {"total_courses": 1, "total_students": 5, "overall_pass_rate": 60,
                   "overall_avg_grade": 70, "highest_avg_course": "Course A", "lowest_avg_course": "Course A"}
        summary = [{"metric": key, "value": value} for key, value in metrics.items()]
        cursor = unittest.mock.MagicMock()
        cursor.fetchall.return_value = [("Course A",)]
        def data(cur, title, sheet, columns, check):
            return [grade] if sheet == "Grade_Distribution" else summary
        with patch.object(m, "google_sheet_records", side_effect=data):
            self.assertEqual(self.run_check(m, m.check_submitted_sheet, cursor), 0)
            grade["pass_rate_pct"] = 100
            self.assertGreater(self.run_check(m, m.check_submitted_sheet, cursor), 0)

    def test_portfolio_report_checks_google_sheet_data_without_xlsx(self):
        m = grader("yf-portfolio-gsheet-pdf-gcal-email")
        symbols = ["GOOGL", "AMZN", "JPM", "JNJ", "XOM"]
        tables = {
            "Holdings": [{"symbol": symbol, "current_price": 100, "shares_held": 100,
                          "market_value": 10000, "allocation_pct": 20} for symbol in symbols],
            "Performance": [{"symbol": symbol, "purchase_price": 90, "current_price": 100,
                             "return_pct": 11.11, "status": "Gain"} for symbol in symbols],
            "Rebalancing": [{"symbol": symbol, "current_allocation": 20, "target_allocation": 20,
                             "drift_pct": 0, "action": "Hold"} for symbol in symbols],
        }
        cursor = unittest.mock.MagicMock()
        cursor.fetchall.return_value = [(symbol, {"currentPrice": 100}) for symbol in symbols]
        with patch.object(m, "google_sheet_records", side_effect=lambda cur, title, sheet, columns, check: tables[sheet]):
            self.assertEqual(self.run_check(m, m.check_submitted_sheet, cursor), 0)
            tables["Performance"][0]["return_pct"] = -20
            self.assertGreater(self.run_check(m, m.check_submitted_sheet, cursor), 0)

    def test_shared_page_title_fix_preserves_content_checks(self):
        for task, title, content in [
            ("fetch-notion-monitoring", "Service Monitoring Dashboard", "Service status report"),
            ("notion-fetch-competitor", "Competitor Overview", "10 products; average price 289; Gamma Ultra"),
        ]:
            for good_title in (True, False):
                with self.subTest(task=task, good_title=good_title):
                    m = grader(task)
                    cursor = unittest.mock.MagicMock()
                    title_value = title if good_title else "Unrelated Notes"
                    cursor.fetchall.side_effect = [
                        [("page-id", {"title": {"title": [{"text": {"content": title_value}}]}})],
                        [("paragraph", {"paragraph": {"rich_text": [{"text": {"content": content}}]}})],
                    ]
                    conn = SimpleNamespace(cursor=lambda: cursor, close=lambda: None)
                    with patch.object(m.psycopg2, "connect", return_value=conn), contextlib.redirect_stdout(io.StringIO()):
                        self.assertEqual(m.check_notion(), good_title)

    def test_customer_crm_title_input_form_keeps_page_requirement(self):
        m = grader("wc-customer-lifetime-excel-notion-email")
        for count in (45, 0):
            cursor = unittest.mock.MagicMock()
            cursor.fetchall.side_effect = [[("crm-db", [{"text": {"content": "Customer CRM"}}], {})],
                                          [(str(i), {}) for i in range(count)]]
            conn = SimpleNamespace(cursor=lambda: cursor, close=lambda: None)
            m.DB_PASS = m.DB_FAIL = 0
            with patch.object(m.psycopg2, "connect", return_value=conn), contextlib.redirect_stdout(io.StringIO()):
                m.check_notion()
            self.assertEqual(m.DB_FAIL, 0 if count == 45 else 1)

    def test_health_query_uses_fixed_task_clock(self):
        m = grader("sf-customer-health-dashboard")
        cursor = unittest.mock.MagicMock()
        cursor.fetchall.return_value = []
        conn = SimpleNamespace(cursor=lambda: cursor, close=lambda: None)
        with patch.object(m.psycopg2, "connect", return_value=conn):
            m.compute_expected_values("2026-03-08 12:30:00")
        sql, args = cursor.execute.call_args.args
        self.assertNotIn("CURRENT_DATE", sql)
        self.assertEqual([v.isoformat() for v in args], ["2026-03-08", "2026-03-08"])


@unittest.skipUnless(os.environ.get("TOOLATHLON_TEST_PG_DSN"), "Requires disposable PostgreSQL database")
class EpochSeededRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import gzip
        with gzip.open(ROOT / "db/init.sql.gz", "rt") as source:
            cls.seed = source.read()

    def setUp(self):
        import psycopg2
        self.conn = psycopg2.connect(os.environ["TOOLATHLON_TEST_PG_DSN"])
        self.conn.set_client_encoding("UTF8")
        self.addCleanup(self.conn.close)
        self.grader_conn = SimpleNamespace(cursor=self.conn.cursor, close=lambda: None)
        self.schemas = set()

    def table(self, name):
        import re
        schema = name.split(".")[0]
        with self.conn.cursor() as cur:
            if schema not in self.schemas:
                cur.execute(f"CREATE SCHEMA {schema}")
                self.schemas.add(schema)
            ddl = re.search(rf"CREATE TABLE {re.escape(name)} \(.*?\);", self.seed, re.S)
            cur.execute(ddl.group())
            copy = re.search(rf"(COPY {re.escape(name)} .*? FROM stdin;)\n(.*?)\n\\\.", self.seed, re.S)
            cur.copy_expert(copy[1], io.StringIO(copy[2] + "\n"))

    def test_google_sheet_reader_uses_actual_seed_schema(self):
        from grader_helpers import google_sheet_records
        for table in ["gsheet.spreadsheets", "gsheet.sheets", "gsheet.cells"]:
            self.table(table)
        with self.conn.cursor() as cur:
            cur.execute("INSERT INTO gsheet.spreadsheets (id, title) VALUES ('epoch-report', 'Report')")
            cur.execute("INSERT INTO gsheet.sheets (id, spreadsheet_id, title) VALUES (901, 'epoch-report', 'Values')")
            for row, values in enumerate([["Metric", "Value"], ["Count", "42"]]):
                for col, value in enumerate(values):
                    cur.execute("INSERT INTO gsheet.cells (spreadsheet_id, sheet_id, row_index, col_index, value) VALUES ('epoch-report', 901, %s, %s, %s)", (row, col, value))
            checks = []
            rows = google_sheet_records(cur, "Report", "Values", ("metric", "value"), lambda name, ok, *rest: checks.append(ok))
        self.assertEqual(rows, [{"metric": "Count", "value": "42"}])
        self.assertTrue(all(checks))

    def test_assignment_expectations_come_from_course_16_seed(self):
        self.table("canvas.assignments")
        self.table("canvas.submissions")
        m = grader("canvas-assignment-feedback-gform-word")
        with patch.object(m.psycopg2, "connect", return_value=self.grader_conn):
            expected = m.expected_assignments()
        names = {row[0] for row in expected}
        self.assertIn("CMA 34878", names)
        self.assertNotIn("CMA 34905", names)
        self.assertEqual(len(expected), 13)
        exam = next(row for row in expected if row[0] == "Final Exam 34885")
        self.assertEqual(exam[1:], (0, None, 0, 0))
        self.assertTrue(all(row[1] > 0 and row[2] > 0 for row in expected if row[0] != exam[0]))

    def test_health_scores_follow_launch_clock_against_real_seed(self):
        for table in ['sf_data."SALES_DW__PUBLIC__ORDERS"', 'sf_data."SALES_DW__PUBLIC__CUSTOMERS"']:
            self.table(table)
        m = grader("sf-customer-health-dashboard")
        with patch.object(m.psycopg2, "connect", return_value=self.grader_conn):
            march = m.compute_expected_values("2026-03-08 12:30:00")
            september = m.compute_expected_values("2026-09-21 12:30:00")
        self.assertEqual(march["total_customers"], september["total_customers"])
        self.assertLess(march["critical_count"], september["critical_count"])
        self.assertEqual(march["critical_count"], 872)

    def test_wc_benchmark_grader_uses_actual_product_prices(self):
        self.table("wc.products")
        m = grader("pw-sf-wc-reconciliation-excel-word")
        with self.conn.cursor() as cur:
            cur.execute("SELECT category->>'name', COUNT(*), AVG(p.price) FROM wc.products p CROSS JOIN LATERAL jsonb_array_elements(p.categories) category GROUP BY category->>'name'")
            products = {name: (count, float(avg)) for name, count, avg in cur.fetchall()}
        benchmarks = m.fixture_table(ROOT / "tasks/finalpool/pw-sf-wc-reconciliation-excel-word")
        rows = [["Category", "Product_Count", "Our_Avg_Price", "Market_Avg_Price", "Price_Gap_Pct"]]
        for ref in sorted(benchmarks, key=lambda row: row["Product_Category"]):
            name = ref["Product_Category"]
            count, average = products[name]
            market = float(ref["Market_Avg_Price"])
            rows.append([name, count, round(average, 2), market, round((average / market - 1) * 100, 2)])
        book = workbook({"Data_Analysis": rows})
        with patch.object(m, "get_conn", return_value=self.grader_conn), contextlib.redirect_stdout(io.StringIO()):
            m.check_benchmark_data(book)
            self.assertEqual(m.FAIL_COUNT, 0)
            book["Data_Analysis"]["D2"] = 100000
            m.check_benchmark_data(book)
            self.assertGreater(m.FAIL_COUNT, 0)

    def test_sector_benchmark_grader_uses_actual_stock_info(self):
        self.table("yf.stock_info")
        m = grader("pw-yf-sector-analysis-excel-word")
        with self.conn.cursor() as cur:
            cur.execute("SELECT data FROM yf.stock_info")
            data = [r[0] for r in cur.fetchall()]
        lookup = {d["sector"]: d for d in data if d.get("sector")}
        benchmarks = m.fixture_table(ROOT / "tasks/finalpool/pw-yf-sector-analysis-excel-word")
        rows = [["Sector", "Our_Avg_PE", "Benchmark_PE", "PE_Gap", "Our_Avg_Yield_Pct", "Benchmark_Yield_Pct", "Yield_Gap_Pct"]]
        for ref in sorted(benchmarks, key=lambda row: row["Sector"]):
            name = ref["Sector"]
            pe, yld = lookup[name]["trailingPE"], lookup[name].get("dividendYield") or 0
            bp, by = float(ref["Benchmark_PE"]), float(ref["Benchmark_Yield_Pct"])
            rows.append([name, pe, bp, pe-bp, yld, by, yld-by])
        book = workbook({"Data_Analysis": rows})
        with patch.object(m, "get_conn", return_value=self.grader_conn), contextlib.redirect_stdout(io.StringIO()):
            m.check_benchmark_data(book)
            self.assertEqual(m.FAIL_COUNT, 0)
            book["Data_Analysis"]["C2"] = 100000
            m.check_benchmark_data(book)
            self.assertGreater(m.FAIL_COUNT, 0)


if __name__ == "__main__":
    unittest.main()
