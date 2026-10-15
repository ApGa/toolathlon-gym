"""No grader may answer a failure to read a deliverable by crediting it.

The c1e58af review found PDF graders that, without PyPDF2 or pdfplumber, only
checked that the PDF existed. A scan of every grader found the same shape around
python-docx, python-pptx and file parsing: an ImportError or exception handler
that measured the file's size, printed PASS, counted a pass or returned success.
This test reads every grader's source and fails on that shape, so a fallback
has to record a failed check instead. Three graders took that path for any file
python-docx or python-pptx could not open; their checks are exercised directly.
Returning success after recording the failure is the same shape: a grader that
gates on the return value credits the file, and one that passes at a share of
its checks (70%) skips the content checks the file would have failed.
"""
import ast
import contextlib
import importlib.util
import io
from pathlib import Path
import random
import shutil
import sys
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
import uuid

TASKS = Path(__file__).resolve().parents[1] / "tasks" / "finalpool"

# Calls that open an agent deliverable, and the modules that read them.
PARSERS = {"Document", "Presentation", "load_workbook", "PdfReader", "pdf_text"}
READER_MODULES = {"docx", "pptx", "openpyxl", "pypdf", "PyPDF2", "pdfplumber", "fitz"}
CAUGHT = {"ImportError", "ModuleNotFoundError", "Exception", "BaseException", None}
CHECK_HELPERS = {"check", "record"}


def _name(node):
    func = node.func
    return func.id if isinstance(func, ast.Name) else getattr(func, "attr", "")


def _caught(handler):
    types = handler.type.elts if isinstance(handler.type, ast.Tuple) else [handler.type]
    return {None if t is None else getattr(t, "id", getattr(t, "attr", None)) for t in types}


def _opens_deliverable(body):
    for node in ast.walk(ast.Module(body=body, type_ignores=[])):
        if isinstance(node, ast.Call):
            if _name(node) in PARSERS:
                return True
            if _name(node) == "open" and getattr(node.func, "value", None) is not None \
                    and getattr(node.func.value, "id", "") in {"pdfplumber", "fitz"}:
                return True
        if isinstance(node, ast.Import) and any(a.name.split(".")[0] in READER_MODULES for a in node.names):
            return True
        if isinstance(node, ast.ImportFrom) and (node.module or "").split(".")[0] in READER_MODULES:
            return True
    return False


def _credits(body):
    """What in a handler stands in for the check it guarded (empty when nothing does)."""
    found, records_failure = [], False
    for node in ast.walk(ast.Module(body=body, type_ignores=[])):
        if isinstance(node, ast.Call):
            name = _name(node)
            if name in {"getsize", "stat"}:
                found.append(f"{name}()")
            if name in CHECK_HELPERS and len(node.args) > 1:
                if isinstance(node.args[1], ast.Constant) and node.args[1].value is True:
                    found.append(f"{name}(..., True)")
                else:
                    records_failure = True
            if name in {"append", "extend"}:
                records_failure = True
            if name == "print":
                for arg in node.args:
                    for sub in ast.walk(arg):
                        if isinstance(sub, ast.Constant) and isinstance(sub.value, str):
                            if "PASS" in sub.value:
                                found.append("prints PASS")
                            if "FAIL" in sub.value:
                                records_failure = True
            if name == "decode" and node.args and getattr(node.args[0], "value", None) in {"latin-1", "latin1"}:
                found.append("decodes raw bytes")
        if isinstance(node, ast.Attribute) and node.attr == "st_size":
            found.append("st_size")
        if isinstance(node, ast.AugAssign) and isinstance(node.op, ast.Add):
            target = getattr(node.target, "id", getattr(node.target, "attr", ""))
            if "pass" in target.lower():
                found.append(f"{target} += 1")
            if "fail" in target.lower():
                records_failure = True
    if _returns_success(body):
        found.append("returns success after recording a failure" if records_failure else "returns success")
    return found


def _returns_success(body):
    for node in ast.walk(ast.Module(body=body, type_ignores=[])):
        if isinstance(node, ast.Return) and node.value is not None:
            value = node.value.elts[0] if isinstance(node.value, ast.Tuple) and node.value.elts else node.value
            if isinstance(value, ast.Constant) and value.value is True:
                return True
    return False


def weakening_handlers(source, filename="<grader>"):
    """(line, what) for each handler around reading a deliverable that credits it instead of failing.

    Also `if <reader> is None:` blocks, where <reader> is what an ImportError handler set to None
    (`try: from pptx import Presentation / except ImportError: Presentation = None`).
    """
    tree = ast.parse(source, filename=filename)
    hits, readers = [], set()
    for node in ast.walk(tree):
        if not isinstance(node, ast.Try):
            continue
        for handler in node.handlers:
            if _caught(handler) & {"ImportError", "ModuleNotFoundError"}:
                readers |= {target.id for stmt in handler.body if isinstance(stmt, ast.Assign)
                            and isinstance(stmt.value, ast.Constant) and stmt.value.value is None
                            for target in stmt.targets if isinstance(target, ast.Name)}
            if not _opens_deliverable(node.body) or not _caught(handler) & CAUGHT:
                continue
            found = _credits(handler.body)
            if found:
                hits.append((handler.lineno, ", ".join(sorted(set(found)))))
    for node in ast.walk(tree):
        test = getattr(node, "test", None)
        if isinstance(node, ast.If) and isinstance(test, ast.Compare) and isinstance(test.left, ast.Name) \
                and test.left.id in readers and len(test.ops) == 1 and isinstance(test.ops[0], ast.Is):
            found = _credits(node.body)
            if found:
                hits.append((node.lineno, ", ".join(sorted(set(found)))))
    return sorted(hits)


class GraderFallbackTest(unittest.TestCase):
    def test_no_grader_credits_a_deliverable_it_could_not_read(self):
        graders = sorted(TASKS.glob("*/evaluation/*.py"))
        self.assertGreater(len(graders), 500)
        hits = []
        for path in graders:
            for line, what in weakening_handlers(path.read_text(), str(path)):
                hits.append(f"{path.relative_to(TASKS)}:{line}: {what}")
        self.assertEqual(hits, [], "fallbacks that credit an unread deliverable:\n" + "\n".join(hits))

    def test_no_grader_reads_a_pdf_as_raw_bytes(self):
        # Searching undecoded bytes credits any file that merely contains the words.
        raw = [str(path.relative_to(TASKS)) for path in sorted(TASKS.glob("*/evaluation/*.py"))
               if 'decode("latin-1"' in path.read_text() or "decode('latin-1'" in path.read_text()]
        self.assertEqual(raw, [])

    def test_the_scan_finds_the_old_shapes(self):
        old_pdf = (
            "def check_pdf(path):\n"
            "    try:\n"
            "        import PyPDF2\n"
            "        text = PyPDF2.PdfReader(path).pages[0].extract_text()\n"
            "    except ImportError:\n"
            "        print('  [WARN] No PDF reader available. Checking existence only.')\n"
            "        return True\n"
        )
        old_docx = (
            "def check_word(path):\n"
            "    try:\n"
            "        from docx import Document\n"
            "        doc = Document(path)\n"
            "    except ImportError:\n"
            "        size = os.path.getsize(path)\n"
            "        record('Word file has content', size > 3000)\n"
            "        return size > 3000\n"
            "    except Exception as e:\n"
            "        record('Word file readable', False, str(e))\n"
            "        return False\n"
        )
        old_pptx = (
            "def check_ppt(path):\n"
            "    try:\n"
            "        from pptx import Presentation\n"
            "        return len(Presentation(path).slides) >= 7, ''\n"
            "    except Exception:\n"
            "        return True, ''\n"
        )
        old_xlsx = (
            "def check_xlsx(path):\n"
            "    try:\n"
            "        wb = openpyxl.load_workbook(path)\n"
            "    except Exception as e:\n"
            "        record('xlsx readable', False, str(e))\n"
            "        return True\n"
        )
        old_none = (
            "try:\n"
            "    from pptx import Presentation\n"
            "except ImportError:\n"
            "    Presentation = None\n"
            "def check_pptx(path):\n"
            "    if Presentation is None:\n"
            "        record('python-pptx available', False, 'Cannot import pptx')\n"
            "        return True  # File exists, can't verify content\n"
        )
        fixed = (
            "try:\n"
            "    from docx import Document\n"
            "except ImportError:\n"
            "    Document = None\n"
            "def check_pdf(path):\n"
            "    if Document is None:\n"
            "        record('python-docx available', False, 'python-docx is not installed')\n"
            "        return False\n"
            "    try:\n"
            "        text = pdf_text(path)\n"
            "    except PdfTextError as error:\n"
            "        check('PDF text readable', False, str(error))\n"
            "        text = ''\n"
            "    try:\n"
            "        wb = openpyxl.load_workbook(path)\n"
            "    except Exception as e:\n"
            "        record('xlsx readable', False, str(e))\n"
            "        wb = None\n"
        )
        self.assertEqual(weakening_handlers(old_pdf), [(5, "returns success")])
        self.assertEqual(weakening_handlers(old_docx), [(5, "getsize()")])
        self.assertEqual(weakening_handlers(old_pptx), [(5, "returns success")])
        self.assertEqual(weakening_handlers(old_xlsx), [(4, "returns success after recording a failure")])
        self.assertEqual(weakening_handlers(old_none), [(6, "returns success after recording a failure")])
        self.assertEqual(weakening_handlers(fixed), [])


def load_grader(task):
    spec = importlib.util.spec_from_file_location(f"grader_{uuid.uuid4().hex}", TASKS / task / "evaluation" / "main.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def presentation(path, slides):
    from pptx import Presentation
    deck = Presentation()
    for i in range(slides):
        deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = f"Coupon {i}"
    deck.save(path)


def document(path, text):
    from docx import Document
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)


class OfficeDeliverableTest(unittest.TestCase):
    """These checks returned success when the Office library raised on the deliverable."""

    CASES = [
        # task, function, file name, writer of a good file, library module, package
        ("wc-coupon-performance-notion-ppt", "check_ppt_file", "Coupon_Analysis_Presentation.pptx",
         lambda path: presentation(path, 8), "pptx", "python-pptx"),
        ("wc-low-stock-reorder-gcal-email", "check_word_file", "Reorder_Report.docx",
         lambda path: document(path, "Low stock reorder report"), "docx", "python-docx"),
        ("wc-order-status-gcal-word", "check_word_file", "Operations_Summary.docx",
         lambda path: document(path, "Order status summary for operations"), "docx", "python-docx"),
    ]

    def test_only_a_readable_deliverable_passes(self):
        for task, function, name, write_good, module, package in self.CASES:
            check = getattr(load_grader(task), function)
            with self.subTest(task=task), TemporaryDirectory() as tmp:
                path = Path(tmp) / name
                write_good(path)
                self.assertEqual(check(tmp), (True, ""))
                with patch.dict(sys.modules, {module: None}):
                    ok, detail = check(tmp)
                self.assertFalse(ok)
                self.assertIn(f"{package} is not installed", detail)
                path.write_bytes(random.Random(7).randbytes(4096))
                ok, detail = check(tmp)
                self.assertFalse(ok)
                self.assertIn("is not a readable", detail)


class UnreadableWorkbookTest(unittest.TestCase):
    """arxiv-research-workflow-pipeline passes at 70% of its checks. It returned on an unreadable workbook,
    leaving 4 checks: workbook and docx exist, docx has a paragraph, workbook unreadable. So random bytes
    named paper_analysis.xlsx and a one-paragraph docx scored 3/4 and passed."""

    TASK = TASKS / "arxiv-research-workflow-pipeline"

    def grade(self, workspace):
        grader = load_grader("arxiv-research-workflow-pipeline")
        argv = ["main.py", "--agent_workspace", str(workspace), "--groundtruth_workspace",
                str(self.TASK / "groundtruth_workspace")]
        out = io.StringIO()
        with patch.object(sys, "argv", argv), contextlib.redirect_stdout(out), self.assertRaises(SystemExit) as code:
            grader.main()
        return code.exception.code, out.getvalue()

    def test_an_unreadable_workbook_fails_every_groundtruth_check(self):
        with TemporaryDirectory() as tmp:
            workspace = Path(tmp)
            document(workspace / "literature_review.docx", "A literature review.")
            (workspace / "paper_analysis.xlsx").write_bytes(random.Random(7).randbytes(8192))
            code, output = self.grade(workspace)
        self.assertEqual(code, 1, output)
        self.assertIn("[FAIL] xlsx readable", output)
        self.assertIn("exists in agent", output)

    def test_the_groundtruth_files_pass(self):
        with TemporaryDirectory() as tmp:
            for name in ("paper_analysis.xlsx", "literature_review.docx"):
                shutil.copy(self.TASK / "groundtruth_workspace" / name, Path(tmp) / name)
            code, output = self.grade(Path(tmp))
        self.assertEqual(code, 0, output)


if __name__ == "__main__":
    unittest.main()
