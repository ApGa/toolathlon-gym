import importlib.util
import subprocess
import sys
from tempfile import TemporaryDirectory
from pathlib import Path
import unittest

import openpyxl

ROOT = Path(__file__).resolve().parents[1] / 'tasks/finalpool'


class HRSummaryGradingTest(unittest.TestCase):
    def grader(self, task):
        spec = importlib.util.spec_from_file_location('hr_summary', ROOT / task / 'evaluation/main.py')
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module

    def test_highest_salary_increase_is_not_first_row_sorted_by_percentage(self):
        task = 'sf-hr-salary-growth'
        wb = openpyxl.load_workbook(ROOT / task / 'groundtruth_workspace/HR_Salary_Growth.xlsx', data_only=True)
        self.addCleanup(wb.close)
        self.assertEqual(wb['Salary Growth']['A2'].value, 'Operations')
        self.assertEqual(self.grader(task).highest_department_names(wb), {'hr'})

    def test_rounded_satisfaction_tie_accepts_either_top_department(self):
        task = 'sf-hr-satisfaction-analysis'
        wb = openpyxl.load_workbook(ROOT / task / 'groundtruth_workspace/HR_Satisfaction_Report.xlsx', data_only=True)
        self.addCleanup(wb.close)
        self.assertEqual(self.grader(task).highest_department_names(wb), {'finance', 'sales'})


    def test_real_salary_report_accepts_computed_highest_and_rejects_stale_summary(self):
        task = ROOT / 'sf-hr-salary-growth'
        with TemporaryDirectory() as tmp:
            workbook = openpyxl.load_workbook(task / 'groundtruth_workspace/HR_Salary_Growth.xlsx')
            self.addCleanup(workbook.close)
            path = Path(tmp) / 'HR_Salary_Growth.xlsx'
            for department, expected_code in [('HR', 0), ('Operations', 1)]:
                workbook['Summary']['B3'] = department
                workbook.save(path)
                result = subprocess.run([sys.executable, str(task / 'evaluation/main.py'),
                                         '--agent_workspace', tmp], capture_output=True, text=True)
                self.assertEqual(result.returncode, expected_code, result.stdout + result.stderr)


if __name__ == '__main__':
    unittest.main()
