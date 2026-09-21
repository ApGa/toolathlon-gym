"""Validate submitted data against the episode's sources, not example workbooks."""
import argparse
import json
import os
import sys
from pathlib import Path

import openpyxl
from grader_helpers import close, number, records, rich_text, text

DB_CONFIG = {"host": os.environ.get("PGHOST", "localhost"), "port": 5432,
             "dbname": os.environ.get("PGDATABASE", "toolathlon_gym"),
             "user": "eigent", "password": "camel"}
PASS_COUNT = FAIL_COUNT = 0


def check(name, condition, detail=""):
    global PASS_COUNT, FAIL_COUNT
    if condition:
        PASS_COUNT += 1
        print(f"  [PASS] {name}")
    else:
        FAIL_COUNT += 1
        print(f"  [FAIL] {name}: {str(detail)[:200]}")


def get_conn():
    import psycopg2
    return psycopg2.connect(**DB_CONFIG)


def check_artifacts(workspace, script, json_files):
    check(f"{script} exists", (workspace / script).is_file())
    for name in json_files:
        try:
            value = json.loads((workspace / name).read_text())
            check(f"{name} contains data", bool(value))
        except (OSError, ValueError) as error:
            check(f"{name} is valid JSON", False, error)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--agent_workspace", default=".")
    parser.add_argument("--groundtruth_workspace", default=".")
    parser.add_argument("--launch_time", default="2026-03-07 10:00:00")
    parser.add_argument("--res_log_file")
    args = parser.parse_args()
    success, message = run_evaluation(args.agent_workspace, args.groundtruth_workspace,
                                      args.launch_time, args.res_log_file)
    print(message)
    sys.exit(0 if success else 1)


def selected_symbols(value):
    """Read normal lists of stock-info objects or mappings keyed by ticker."""
    if isinstance(value, list):
        return {symbol for item in value for symbol in selected_symbols(item)}
    if isinstance(value, dict):
        direct = value.get("symbol", value.get("ticker"))
        if direct is not None:
            return {str(direct).upper()}
        found = set()
        for key, item in value.items():
            if isinstance(item, dict) and ("currentPrice" in item or "marketCap" in item):
                found.add(str(key).upper())
            else:
                found.update(selected_symbols(item))
        return found
    return set()


def check_sectors(book, selected, stocks, papers):
    sectors = records(book, "Sector_Performance", ("sector", "stock_count", "avg_price", "total_market_value", "volatility_score"), check)
    research = records(book, "Research_Mapping", ("paper_title", "key_finding", "applicable_sector", "validation_status"), check)
    thesis = records(book, "Investment_Thesis", ("sector", "outlook", "supporting_evidence", "risk_factor"), check)
    lookup = {symbol: data for symbol, data in stocks}
    check("Selected tickers exist in financial source", bool(selected) and selected <= set(lookup))
    expected = {}
    for symbol in selected & set(lookup):
        data = lookup[symbol]
        price = number(data.get("currentPrice") if data.get("currentPrice") is not None
                       else data.get("regularMarketPrice"))
        market_cap = number(data.get("marketCap"))
        valid = bool(data.get("sector")) and price is not None and market_cap is not None
        check(f"Selected stock has sector, price and market capitalization: {symbol}", valid)
        if not valid:
            continue
        values = expected.setdefault(text(data["sector"]), [0, 0.0, 0.0])
        values[0] += 1
        values[1] += price
        values[2] += market_cap
    names = [text(row["sector"]) for row in sectors]
    check("Every selected sector appears once, alphabetically", len(expected) >= 2 and len(names) == len(expected)
          and set(names) == set(expected) and names == sorted(names))
    for row in sectors:
        values = expected.get(text(row["sector"]))
        if values is None:
            continue
        count, prices, market_cap = values
        check(f"Financial source totals: {row['sector']}", close(row["stock_count"], count, 0)
              and close(row["avg_price"], prices / count, .02)
              and close(row["total_market_value"], market_cap, .02))
        # The task does not define a volatility estimator. It must be reported,
        # but a placeholder workbook is not a unique financial answer.
        volatility = number(row["volatility_score"])
        check("Volatility score is a nonnegative number", volatility is not None and volatility >= 0)
    titles = [text(row["paper_title"]) for row in research]
    available = {text(title) for title in papers}
    check("Research uses distinct available papers, sorted by title", bool(titles)
          and len(titles) == len(set(titles)) and set(titles) <= available and titles == sorted(titles))
    for row in research:
        check("Research finding and validation verdict present", bool(text(row["key_finding"]))
              and bool(text(row["applicable_sector"]))
              and text(row["validation_status"]) in {"confirmed", "partial", "inconclusive"})
    thesis_sectors = [text(row["sector"]) for row in thesis]
    check("Investment thesis covers selected sectors once", len(thesis_sectors) == len(expected)
          and set(thesis_sectors) == set(expected))
    for row in thesis:
        check("Investment thesis includes evidence and risk", text(row["outlook"]) in {"bullish", "neutral", "bearish"}
              and bool(text(row["supporting_evidence"])) and bool(text(row["risk_factor"])))


def run_evaluation(agent_workspace, groundtruth_workspace, launch_time, res_log_file):
    global PASS_COUNT, FAIL_COUNT
    PASS_COUNT = FAIL_COUNT = 0
    workspace = Path(agent_workspace)
    check_artifacts(workspace, "sector_analyst.py", ["financial_data.json", "research_findings.json", "sector_analysis.json"])
    path = workspace / "Sector_Analysis_Report.xlsx"
    check(path.name + " exists", path.is_file())
    selected = set()
    try:
        selected = selected_symbols(json.loads((workspace / "financial_data.json").read_text()))
    except (OSError, ValueError):
        pass
    conn = get_conn()
    try:
        with conn.cursor() as cursor:
            cursor.execute("SELECT symbol, data FROM yf.stock_info")
            stocks = cursor.fetchall()
            cursor.execute("SELECT title FROM scholarly.scholar_papers UNION SELECT title FROM scholarly.arxiv_papers")
            papers = [row[0] for row in cursor.fetchall()]
            if path.is_file():
                book = openpyxl.load_workbook(path, data_only=True)
                try:
                    check_sectors(book, selected, stocks, papers)
                finally:
                    book.close()
    finally:
        conn.close()
    doc_path = workspace / "Sector_Research_Brief.docx"
    check(doc_path.name + " exists", doc_path.is_file())
    if doc_path.is_file():
        from docx import Document
        doc = Document(doc_path)
        body = " ".join(p.text for p in doc.paragraphs)
        headings = {text(p.text) for p in doc.paragraphs}
        check("Research brief has content", len(body) > 50)
        for heading in ["Financial Performance Review", "Academic Research Insights", "Theory vs Practice Comparison", "Investment Implications"]:
            check(f"Brief section: {heading}", text(heading) in headings)
    return FAIL_COUNT == 0, f"Passed {PASS_COUNT}/{PASS_COUNT + FAIL_COUNT} checks"


if __name__ == "__main__":
    main()
