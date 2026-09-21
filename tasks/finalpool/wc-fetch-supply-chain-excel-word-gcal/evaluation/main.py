"""Validate submitted data against the episode's sources, not example workbooks."""
import argparse
import json
import os
import sys
from pathlib import Path

import openpyxl
from grader_helpers import close, number, records, rich_text, text

DB_CONFIG = {"host": os.environ.get("PGHOST", "localhost"), "port": 5432,
             "dbname": os.environ.get("PGDATABASE", "toolathlon_gym"),
             "user": "eigent", "password": "camel"}
PASS_COUNT = FAIL_COUNT = 0


def check(name, condition, detail=""):
    global PASS_COUNT, FAIL_COUNT
    if condition:
        PASS_COUNT += 1
        print(f"  [PASS] {name}")
    else:
        FAIL_COUNT += 1
        print(f"  [FAIL] {name}: {str(detail)[:200]}")


def get_conn():
    import psycopg2
    return psycopg2.connect(**DB_CONFIG)


def check_artifacts(workspace, script, json_files):
    check(f"{script} exists", (workspace / script).is_file())
    for name in json_files:
        try:
            value = json.loads((workspace / name).read_text())
            check(f"{name} contains data", bool(value))
        except (OSError, ValueError) as error:
            check(f"{name} is valid JSON", False, error)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--agent_workspace", default=".")
    parser.add_argument("--groundtruth_workspace", default=".")
    parser.add_argument("--launch_time", default="2026-03-07 10:00:00")
    parser.add_argument("--res_log_file")
    args = parser.parse_args()
    success, message = run_evaluation(args.agent_workspace, args.groundtruth_workspace,
                                      args.launch_time, args.res_log_file)
    print(message)
    sys.exit(0 if success else 1)


def source_suppliers():
    import tarfile
    task = Path(__file__).resolve().parents[1]
    with tarfile.open(task / "files/mock_pages.tar.gz") as archive:
        member = next(item for item in archive.getmembers() if item.name.endswith("api/supplier_directory.json"))
        return json.load(archive.extractfile(member))["suppliers"]


def check_supply_chain(book, products, suppliers):
    inventory = records(book, "Inventory_Status", ("product_name", "current_stock", "daily_sales_rate", "days_until_stockout", "reorder_point", "status"), check)
    directory = records(book, "Supplier_Analysis", ("supplier_name", "lead_time_days", "reliability_score", "min_order_qty", "products_supplied"), check)
    plan = records(book, "Reorder_Plan", ("product_name", "current_stock", "reorder_qty", "supplier", "expected_delivery_date", "urgency"), check)
    summary = records(book, "Summary", ("metric", "value"), check)
    stock = {text(name): (quantity, float(price)) for name, quantity, price in products}
    supplier_lookup = {text(row["name"]): row for row in suppliers}
    planned_suppliers = {text(row["product_name"]): text(row["supplier"]) for row in plan}
    names = [text(row["product_name"]) for row in inventory]
    check("All store products represented once", len(names) == len(stock) and set(names) == set(stock))
    days, statuses = [], []
    for row in inventory:
        name = text(row["product_name"])
        if name not in stock:
            continue
        quantity = stock[name][0]
        rate, point = number(row["daily_sales_rate"]), number(row["reorder_point"])
        check(f"Store stock: {name}", close(row["current_stock"], quantity, 0))
        valid = rate is not None and rate >= 0 and point is not None and point >= 0
        check(f"Nonnegative rate/reorder point: {name}", valid)
        if not valid:
            continue
        # No demand window or unique supplier assignment is prescribed. Check
        # the stated 1.5 formula for a real supplier, allowing displayed rounding.
        source_options = ([supplier_lookup[planned_suppliers[name]]]
                          if planned_suppliers.get(name) in supplier_lookup else suppliers)
        candidates = [rate * supplier["lead_time_days"] * 1.5 for supplier in source_options
                      if abs(point - rate * supplier["lead_time_days"] * 1.5)
                      <= .51 + .005 * supplier["lead_time_days"] * 1.5]
        check(f"Reorder formula with available supplier: {name}", bool(candidates))
        expected_statuses = {"critical" if quantity < candidate else "warning" if quantity <= 1.2 * candidate else "healthy"
                             for candidate in candidates + ([point] if candidates else [])}
        status = text(row["status"])
        check(f"Inventory status follows reorder threshold: {name}", status in expected_statuses)
        statuses.append(status)
        actual_days = number(row["days_until_stockout"])
        if rate > 0:
            lower = quantity / (rate + .005)
            upper = quantity / max(rate - .005, 1e-12)
            check(f"Days of stock follow reported rate: {name}", actual_days is not None
                  and lower - .51 <= actual_days <= upper + .51)
            days.append(actual_days if actual_days is not None else float("inf"))
        else:
            check(f"No-demand stock duration marked unavailable: {name}",
                  text(row["days_until_stockout"]) in {"n/a", "inf", "infinity", ""})
            days.append(float("inf"))
    check("Inventory sorted by days until stockout", days == sorted(days))
    supplier_names = [text(row["supplier_name"]) for row in directory]
    check("Supplier directory represented once", len(supplier_names) == len(suppliers)
          and set(supplier_names) == set(supplier_lookup))
    reliability, supplied = [], []
    for row in directory:
        source = supplier_lookup.get(text(row["supplier_name"]))
        if source is None:
            continue
        check(f"Supplier fixture values: {source['name']}", all(close(row[key], source[key], 0)
              for key in ("lead_time_days", "reliability_score", "min_order_qty")))
        reliability.append(source["reliability_score"])
        count = number(row["products_supplied"])
        check("Products supplied is a nonnegative integer", count is not None and count >= 0 and count.is_integer())
        supplied.append(count or 0)
    check("Suppliers sorted by reliability", reliability == sorted(reliability, reverse=True))
    check("Supplier counts cover inventory", sum(supplied) == len(stock))
    order_value = 0.0
    planned = []
    for row in plan:
        name, supplier = text(row["product_name"]), text(row["supplier"])
        valid = name in stock and supplier in supplier_lookup
        check("Reorder uses real product and supplier", valid)
        if not valid:
            continue
        planned.append(name)
        check(f"Reorder stock: {name}", close(row["current_stock"], stock[name][0], 0))
        quantity = number(row["reorder_qty"])
        check("Reorder quantity is a positive integer", quantity is not None and quantity > 0 and quantity.is_integer())
        order_value += (quantity or 0) * stock[name][1]
        from datetime import datetime
        try:
            datetime.fromisoformat(str(row["expected_delivery_date"]))
            date_valid = True
        except ValueError:
            date_valid = False
        check("Delivery date and urgency provided", date_valid and text(row["urgency"]) in {"immediate", "this week", "this month"})
    critical = {text(row["product_name"]) for row in inventory if text(row["status"]) == "critical"}
    check("Reorder plan covers critical products without duplicates", critical <= set(planned) and len(planned) == len(set(planned)))
    metrics = {text(row["metric"]): row["value"] for row in summary}
    for label, expected in [("total_products", len(stock)), ("critical_products", statuses.count("critical")),
                            ("warning_products", statuses.count("warning")), ("healthy_products", statuses.count("healthy")),
                            ("total_reorder_value", order_value)]:
        check(f"Summary {label}", close(metrics.get(label), expected, .02))
    finite_days = [value for value in days if value != float("inf")]
    average_matches = (close(metrics.get("avg_days_until_stockout"), sum(finite_days) / len(finite_days), .56)
                       if finite_days else text(metrics.get("avg_days_until_stockout")) in {"n/a", ""})
    check("Summary average stock duration", average_matches)


def run_evaluation(agent_workspace, groundtruth_workspace, launch_time, res_log_file):
    global PASS_COUNT, FAIL_COUNT
    PASS_COUNT = FAIL_COUNT = 0
    workspace = Path(agent_workspace)
    check_artifacts(workspace, "supply_chain_optimizer.py", ["supplier_data.json", "inventory_data.json", "supply_chain_plan.json"])
    path = workspace / "Supply_Chain_Optimization.xlsx"
    check(path.name + " exists", path.is_file())
    conn = get_conn()
    try:
        with conn.cursor() as cursor:
            cursor.execute("SELECT name, stock_quantity, price FROM wc.products WHERE status = 'publish'")
            products = cursor.fetchall()
            if path.is_file():
                book = openpyxl.load_workbook(path, data_only=True)
                try:
                    check_supply_chain(book, products, source_suppliers())
                finally:
                    book.close()
            cursor.execute("""SELECT COUNT(*) FROM gcal.events WHERE summary = 'Supply Chain Review Meeting'
                              AND start_datetime = '2026-03-19 09:00:00+00'::timestamptz
                              AND end_datetime = '2026-03-19 10:30:00+00'::timestamptz""")
            check("Supply Chain Review Meeting at requested time", cursor.fetchone()[0] >= 1)
            cursor.execute("SELECT COUNT(*) FROM gcal.events WHERE summary IN ('Team Standup', 'Lunch Break')")
            check("Existing calendar events preserved", cursor.fetchone()[0] >= 1)
    finally:
        conn.close()
    doc_path = workspace / "Supply_Chain_Report.docx"
    check(doc_path.name + " exists", doc_path.is_file())
    if doc_path.is_file():
        from docx import Document
        doc = Document(doc_path)
        body = " ".join(p.text for p in doc.paragraphs)
        headings = {text(p.text) for p in doc.paragraphs}
        check("Supply-chain report has content", len(body) > 50)
        for heading in ["Inventory Health Assessment", "Supplier Performance Review", "Reorder Recommendations", "Risk Mitigation Plan"]:
            check(f"Report section: {heading}", text(heading) in headings)
    return FAIL_COUNT == 0, f"Passed {PASS_COUNT}/{PASS_COUNT + FAIL_COUNT} checks"


if __name__ == "__main__":
    main()
