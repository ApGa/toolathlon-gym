"""Evaluation script for pw-sf-wc-reconciliation-excel-word."""
import os
import argparse, json, os, sys
import openpyxl
from pathlib import Path
from grader_helpers import records, close, text, fixture_table


DB_CONFIG = {
    "host": os.environ.get("PGHOST", "localhost"), "port": 5432,
    "dbname": os.environ.get("PGDATABASE", "toolathlon_gym"),
    "user": "eigent", "password": "camel"
}

PASS_COUNT = 0
FAIL_COUNT = 0

def check(name, condition, detail=""):
    global PASS_COUNT, FAIL_COUNT
    if condition:
        PASS_COUNT += 1
        print(f"  [PASS] {name}")
    else:
        FAIL_COUNT += 1
        detail_str = str(detail)[:200] if detail else ""
        print(f"  [FAIL] {name}: {detail_str}")

def safe_float(val, default=None):
    try:
        if val is None:
            return default
        return float(str(val).replace(',', '').replace('%', '').replace('$', '').strip())
    except (ValueError, TypeError):
        return default

def get_conn():
    import psycopg2
    return psycopg2.connect(**DB_CONFIG)

def check_benchmark_data(wb):
    benchmark = fixture_table(Path(__file__).resolve().parents[1])
    conn = get_conn()
    try:
        with conn.cursor() as cur:
            # Match the store's embedded category names, as exposed by product
            # tools, instead of an unrelated sales-region example workbook.
            cur.execute("""
                SELECT category->>'name', COUNT(*), AVG(p.price)
                FROM wc.products p CROSS JOIN LATERAL jsonb_array_elements(p.categories) category
                GROUP BY category->>'name'
            """)
            products = {name: (count, float(avg)) for name, count, avg in cur.fetchall()}
    finally:
        conn.close()
    rows = records(wb, "Data_Analysis", ("category", "product_count", "our_avg_price", "market_avg_price", "price_gap_pct"), check)
    lookup = {text(r["category"]): r for r in rows}
    check("One row for every benchmark category", len(rows) == len(benchmark)
          and set(lookup) == {text(r["Product_Category"]) for r in benchmark})
    for reference in benchmark:
        category = reference["Product_Category"]
        count, average = products.get(category, (0, 0))
        market = float(reference["Market_Avg_Price"])
        row = lookup.get(text(category), {})
        check(f"{category} comparison values", count > 0 and close(row.get("product_count"), count, 0)
              and close(row.get("our_avg_price"), average, .02)
              and close(row.get("market_avg_price"), market, .02)
              and close(row.get("price_gap_pct"), (average / market - 1) * 100, .02))
    dimensions = [text(r["category"]) for r in rows]
    check("Categories sorted alphabetically", dimensions == sorted(dimensions))


def run_evaluation(agent_workspace, groundtruth_workspace, launch_time, res_log_file):
    global PASS_COUNT, FAIL_COUNT
    PASS_COUNT = 0
    FAIL_COUNT = 0

    
    excel_path = os.path.join(agent_workspace, "Wc_Reconciliation_Report.xlsx")
    check("Wc_Reconciliation_Report.xlsx exists", os.path.exists(excel_path))
    if os.path.exists(excel_path):
        wb = openpyxl.load_workbook(excel_path)
        gt_path = os.path.join(groundtruth_workspace, "Wc_Reconciliation_Report.xlsx")
        gt_wb = openpyxl.load_workbook(gt_path) if os.path.exists(gt_path) else None

        check_benchmark_data(wb)

        check("Metrics sheet exists", "Metrics" in wb.sheetnames)
        if "Metrics" in wb.sheetnames:
            ws = wb["Metrics"]
            data_rows = list(ws.iter_rows(min_row=2, values_only=True))
            check("Metrics has >= 4 rows", len(data_rows) >= 4, f"got {len(data_rows)}")

            # Check headers
            headers = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
            for expected_col in ['Metric', 'Value']:
                check(f"Metrics has {expected_col} column",
                      expected_col.lower() in headers, f"headers: {headers[:8]}")

        check("Recommendations sheet exists", "Recommendations" in wb.sheetnames)
        if "Recommendations" in wb.sheetnames:
            ws = wb["Recommendations"]
            data_rows = list(ws.iter_rows(min_row=2, values_only=True))
            check("Recommendations has >= 2 rows", len(data_rows) >= 2, f"got {len(data_rows)}")

            # Check headers
            headers = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
            for expected_col in ['Priority', 'Action', 'Category']:
                check(f"Recommendations has {expected_col} column",
                      expected_col.lower() in headers, f"headers: {headers[:8]}")

        # Check Word document
        import glob as globmod
        word_files = globmod.glob(os.path.join(agent_workspace, "*.docx"))
        check("Word document exists", len(word_files) >= 1, f"found {len(word_files)} docx files")
        if word_files:
            from docx import Document
            doc = Document(word_files[0])
            text = " ".join(p.text for p in doc.paragraphs).lower()
            check("Word has content", len(text) > 50, f"text length: {len(text)}")

        check("sf_wc_reconcile_processor.py exists", os.path.exists(os.path.join(agent_workspace, "sf_wc_reconcile_processor.py")))


    return FAIL_COUNT == 0, f"Passed {PASS_COUNT}/{PASS_COUNT + FAIL_COUNT} checks"

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--agent_workspace", required=False, default=".")
    parser.add_argument("--groundtruth_workspace", required=False, default=".")
    parser.add_argument("--launch_time", required=False, default="2026-03-07 10:00:00")
    parser.add_argument("--res_log_file", required=False)
    args = parser.parse_args()

    success, message = run_evaluation(
        args.agent_workspace, args.groundtruth_workspace,
        args.launch_time, args.res_log_file
    )
    print(message)
    sys.exit(0 if success else 1)

if __name__ == "__main__":
    main()
